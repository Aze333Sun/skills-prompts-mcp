#!/usr/bin/env python3
"""将方案设计说明书 Markdown 转换为符合中国联通模板样式的 Word 文档"""

import os
import re
import sys
from docx import Document
from docx.shared import Pt, Twips, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_font_full(rPr, font_name='Times New Roman', east_asia='宋体'):
    """完整设置字体：清除主题字体引用，显式设置 ascii/hAnsi/eastAsia/cs 四个属性。
    Word 的 Heading 样式内置 asciiTheme/eastAsiaTheme 等主题字体引用，
    会覆盖直接设置的字体（导致回退到 MS ゴシック / Calibri Light），必须先清除。"""
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    # 清除主题字体属性（使用 qn() 获取带命名空间的全限定名）
    for theme_attr in [qn('w:asciiTheme'), qn('w:hAnsiTheme'), qn('w:eastAsiaTheme'), qn('w:cstheme')]:
        if theme_attr in rFonts.attrib:
            del rFonts.attrib[theme_attr]
    # 显式设置四个字体属性
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), east_asia)
    rFonts.set(qn('w:cs'), font_name)

def fill_cell(cell, text, size, bold, align):
    """填充单元格：支持行内 **加粗**、`代码` 与 <br> 换行，避免按字面输出"""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    lines = re.split(r'<br\s*/?>', text)
    for line_idx, line in enumerate(lines):
        if line_idx > 0:
            paragraph.add_run().add_break()
        for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', line):
            if not part:
                continue
            run = paragraph.add_run()
            if part.startswith('**') and part.endswith('**') and len(part) > 4:
                run.text = part[2:-2]
                run.font.bold = True
                set_font_full(run._element.get_or_add_rPr(), 'Times New Roman', '宋体')
            elif part.startswith('`') and part.endswith('`') and len(part) > 2:
                run.text = part[1:-1]
                set_font_full(run._element.get_or_add_rPr(), 'Courier New', '宋体')
            else:
                run.text = part
                run.font.bold = bold
                set_font_full(run._element.get_or_add_rPr(), 'Times New Roman', '宋体')
            run.font.size = size

def set_run_font(run, font_name='Times New Roman', east_asia='宋体', size=Pt(12), bold=False, color=None):
    run.font.name = font_name
    run.font.size = size
    run.font.bold = bold
    set_font_full(run._element.get_or_add_rPr(), font_name, east_asia)
    if color:
        run.font.color.rgb = color

def clear_theme_color(rPr):
    """清除主题颜色引用（accent1 等），改为自动颜色，避免标题内置蓝色。"""
    color = rPr.find(qn('w:color'))
    if color is not None:
        for attr in list(color.attrib.keys()):
            if attr.startswith('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}themeColor') or \
               attr.startswith('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}themeShade') or \
               attr.startswith('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}themeTint'):
                del color.attrib[attr]
        color.set(qn('w:val'), 'auto')

def setup_styles(doc):
    # 正文 Normal: 中文宋体 数字英文Times New Roman 11pt 不加粗 两端对齐 行距1.5倍
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(11)
    set_font_full(style_normal._element.get_or_add_rPr(), 'Times New Roman', '宋体')
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_normal.paragraph_format.line_spacing = 1.5

    # 标题样式统一处理：清除主题字体和主题颜色（避免内置蓝色和 MS ゴシック）
    for style_name in ['Heading 1', 'Heading 2', 'Heading 3', 'Heading 4']:
        clear_theme_color(doc.styles[style_name]._element.get_or_add_rPr())

    # 标题1 Heading 1: 中文宋体 小二(18pt) 加粗 左对齐 段前1行段后0.5行
    style_h1 = doc.styles['Heading 1']
    style_h1.font.name = 'Times New Roman'
    style_h1.font.size = Pt(18)
    style_h1.font.bold = True
    set_font_full(style_h1._element.get_or_add_rPr(), 'Times New Roman', '宋体')
    style_h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_h1.paragraph_format.space_before = Pt(18)
    style_h1.paragraph_format.space_after = Pt(9)

    # 标题2 Heading 2: 中文宋体 三号(16pt) 加粗 左对齐 段前0.5行段后0.3行
    style_h2 = doc.styles['Heading 2']
    style_h2.font.name = 'Times New Roman'
    style_h2.font.size = Pt(16)
    style_h2.font.bold = True
    set_font_full(style_h2._element.get_or_add_rPr(), 'Times New Roman', '宋体')
    style_h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_h2.paragraph_format.space_before = Pt(10)
    style_h2.paragraph_format.space_after = Pt(6)

    # 标题3 Heading 3: 中文宋体 四号(14pt) 加粗 左对齐 段前0.3行
    style_h3 = doc.styles['Heading 3']
    style_h3.font.name = 'Times New Roman'
    style_h3.font.size = Pt(14)
    style_h3.font.bold = True
    set_font_full(style_h3._element.get_or_add_rPr(), 'Times New Roman', '宋体')
    style_h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_h3.paragraph_format.space_before = Pt(6)

    # 标题4 Heading 4: 中文宋体 小四(12pt) 不加粗 左对齐
    style_h4 = doc.styles['Heading 4']
    style_h4.font.name = 'Times New Roman'
    style_h4.font.size = Pt(12)
    style_h4.font.bold = False
    set_font_full(style_h4._element.get_or_add_rPr(), 'Times New Roman', '宋体')
    style_h4.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def setup_page(doc):
    # A4 纵向，上下边距 2.54cm (1440 twips)，左右边距 3.17cm (1797 twips)
    # 无页眉、无页脚、无页码
    section = doc.sections[0]
    section.page_width = Twips(11906)
    section.page_height = Twips(16838)
    section.top_margin = Twips(1440)
    section.bottom_margin = Twips(1440)
    section.left_margin = Twips(1797)
    section.right_margin = Twips(1797)
    section.gutter = Twips(0)
    section.different_first_page_header_footer = False
    # 清空页眉页脚内容
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for hdr in [section.first_page_header, section.header]:
        for p in hdr.paragraphs:
            p.text = ''
    for ftr in [section.first_page_footer, section.footer]:
        for p in ftr.paragraphs:
            p.text = ''

def resolve_image_path(img_path, bases):
    """在多个候选基目录下解析图片路径，避免因运行目录不同导致误判「图片不存在」。

    bases: 候选基目录列表（如 CWD、Markdown 文件所在目录、项目根目录）。
    返回首个真实存在的路径；全部不存在则返回 None。
    """
    if os.path.isabs(img_path) and os.path.exists(img_path):
        return img_path
    for base in bases:
        if not base:
            continue
        cand = os.path.normpath(os.path.join(base, img_path))
        if os.path.exists(cand):
            return cand
    return None

def parse_markdown_table(lines):
    rows = []
    for line in lines:
        line = line.strip()
        if line.startswith('|') and line.endswith('|'):
            cells = [c.strip() for c in line[1:-1].split('|')]
            rows.append(cells)
    filtered = []
    for row in rows:
        if all(re.match(r'^[-:]+$', c) for c in row):
            continue
        filtered.append(row)
    return filtered

def add_table_to_doc(doc, table_rows):
    if not table_rows:
        return
    num_cols = len(table_rows[0])
    table = doc.add_table(rows=len(table_rows), cols=num_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row_data in enumerate(table_rows):
        for j, cell_text in enumerate(row_data):
            if j < num_cols:
                cell = table.cell(i, j)
                # 表格标题：12pt 加粗 居中；表格文字：10.5pt 不加粗 左对齐
                if i == 0:
                    fill_cell(cell, cell_text, Pt(12), True, WD_ALIGN_PARAGRAPH.CENTER)
                else:
                    fill_cell(cell, cell_text, Pt(10.5), False, WD_ALIGN_PARAGRAPH.LEFT)

def add_paragraph_with_runs(doc, text, style_name=None):
    p = doc.add_paragraph(style=style_name)
    parts = re.split(r'(\*\*.*?\*\*|`[^`]+`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.font.bold = True
            if not style_name or 'Heading' not in style_name:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11)
                set_font_full(run._element.get_or_add_rPr(), 'Times New Roman', '宋体')
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(11)
            set_font_full(run._element.get_or_add_rPr(), 'Courier New', '宋体')
        elif part:
            run = p.add_run(part)
            if not style_name or 'Heading' not in style_name:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11)
                set_font_full(run._element.get_or_add_rPr(), 'Times New Roman', '宋体')
    return p

def convert_md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # 自动创建输出目录（如 方案/），避免因目录不存在导致 doc.save 失败
    out_dir = os.path.dirname(os.path.abspath(docx_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    # 图片路径解析基目录：CWD / Markdown 所在目录 / 项目根（Markdown 通常位于 <root>/.trae/tmp/ 下，向上两级）
    md_dir = os.path.dirname(os.path.abspath(md_path))
    project_root = os.path.dirname(os.path.dirname(md_dir))
    img_bases = [os.getcwd(), md_dir, project_root]

    doc = Document()
    setup_page(doc)
    setup_styles(doc)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # 水平分隔线：跳过，不渲染为字面文本
        if re.match(r'^(-{3,}|\*{3,})$', stripped):
            i += 1
            continue

        if stripped.startswith('```'):
            # 整块收集代码内容，按等宽字体原样输出，避免被当作 Markdown 二次解析
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\r\n'))
                i += 1
            i += 1  # 跳过结束围栏
            for code_line in code_lines:
                p = doc.add_paragraph()
                run = p.add_run(code_line)
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
                set_font_full(run._element.get_or_add_rPr(), 'Courier New', '宋体')
                p.paragraph_format.left_indent = Twips(420)
                p.paragraph_format.line_spacing = 1.0
            continue

        # 图片：嵌入 Word 文档，超出版心宽度时等比缩小
        img_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', stripped)
        if img_match:
            img_path = img_match.group(2).strip()
            resolved = resolve_image_path(img_path, img_bases)
            if resolved:
                doc.add_picture(resolved)
                shape = doc.inline_shapes[-1]
                content_width = Twips(8312)  # 版心宽度 = 页宽11906 - 左右边距1797*2
                if shape.width > content_width:
                    ratio = content_width / shape.width
                    shape.height = int(shape.height * ratio)
                    shape.width = int(content_width)
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                # 图片缺失：输出明确占位提示，避免把原始 markdown 链接当作字面文本写入文档
                p = doc.add_paragraph()
                run = p.add_run(f'[图片缺失：{img_path}]')
                set_font_full(run._element.get_or_add_rPr(), 'Times New Roman', '宋体')
                run.font.size = Pt(11)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if stripped.startswith('#'):
            match = re.match(r'^(#{1,4})\s+(.*)', stripped)
            if match:
                level = len(match.group(1))
                title_text = match.group(2)
                style_map = {1: 'Heading 1', 2: 'Heading 2', 3: 'Heading 3', 4: 'Heading 4'}
                style_name = style_map.get(level, 'Heading 4')
                p = doc.add_paragraph(style=style_name)
                run = p.add_run(title_text)
                # 显式设置 run 字体，清除主题字体引用，避免回退到 MS ゴシック
                set_font_full(run._element.get_or_add_rPr(), 'Times New Roman', '宋体')
                # 清除 run 级别的主题颜色
                clear_theme_color(run._element.get_or_add_rPr())
                i += 1
                continue

        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            table_rows = parse_markdown_table(table_lines)
            add_table_to_doc(doc, table_rows)
            doc.add_paragraph()
            continue

        if re.match(r'^[-*]\s+', stripped):
            content = re.sub(r'^[-*]\s+', '', stripped)
            add_paragraph_with_runs(doc, content)
            i += 1
            continue

        if re.match(r'^\d+\.\s+', stripped):
            # 保留编号文本，避免列表序号丢失
            p = add_paragraph_with_runs(doc, stripped)
            p.paragraph_format.left_indent = Twips(420)
            i += 1
            continue

        # 引用块：去掉 > 前缀后缩进渲染
        if stripped.startswith('>'):
            p = add_paragraph_with_runs(doc, stripped.lstrip('> ').strip())
            p.paragraph_format.left_indent = Twips(420)
            i += 1
            continue

        add_paragraph_with_runs(doc, stripped)
        i += 1

    doc.save(docx_path)
    print(f'Word 文档已生成: {docx_path}')

if __name__ == '__main__':
    import argparse

    parser = argparse.ArgumentParser(
        description='将方案设计说明书 Markdown 转换为符合中国联通模板样式的 Word 文档',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='''
使用示例:
  python3 md2docx.py document.md
  python3 md2docx.py document.md -o output.docx
  python3 md2docx.py /path/to/design.md -o /path/to/result.docx
        '''
    )
    parser.add_argument('input', help='输入的 Markdown 文件路径')
    parser.add_argument('-o', '--output', help='输出的 Word 文件路径（默认与输入同目录，扩展名改为 .docx）')

    args = parser.parse_args()

    md_path = args.input
    if not os.path.exists(md_path):
        print(f'错误：文件不存在 - {md_path}')
        sys.exit(1)

    if args.output:
        docx_path = args.output
    else:
        base, _ = os.path.splitext(md_path)
        docx_path = base + '.docx'

    convert_md_to_docx(md_path, docx_path)
